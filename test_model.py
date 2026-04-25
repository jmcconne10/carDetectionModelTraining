import os
import logging
from datetime import datetime
from ultralytics import YOLO
from openpyxl import load_workbook
from openpyxl.styles import numbers

logging.getLogger("ultralytics").setLevel(logging.WARNING)

test_folder = "newTests"
model_path = "/Users/joemcconnell/Documents/Code/yolo-car-test/best.pt"
tracker_path = "/Users/joemcconnell/Documents/Code/yolo-car-test/model_tracker.xlsx"

if not os.path.exists(test_folder):
    print(f"Folder not found: {test_folder}")
    exit()

if not os.path.exists(model_path):
    print(f"Model not found: {model_path}")
    exit()

if not os.path.exists(tracker_path):
    print(f"Tracker not found: {tracker_path}")
    exit()

wb = load_workbook(tracker_path)
ws = wb.active

expected = {}
for row in range(2, ws.max_row + 1):
    img_name = ws.cell(row=row, column=1).value
    car = ws.cell(row=row, column=2).value
    if img_name and car:
        expected[img_name] = car.strip().lower()

model = YOLO(model_path)

print(f"\n{'File':<40} {'Expected':<12} {'Predicted':<12} {'Result':<12}")
print("-" * 76)

run_results = {}

for img_name in expected:
    img_path = os.path.join(test_folder, img_name)
    if not os.path.exists(img_path):
        print(f"{img_name:<40} {'FILE MISSING'}")
        run_results[img_name] = "missing"
        continue

    results = model(img_path, save=True, conf=0.25, verbose=False)

    if len(results[0].boxes) == 0:
        print(f"{img_name:<40} {expected[img_name]:<12} {'none':<12} {'no detection':<12}")
        run_results[img_name] = "no detection"
    else:
        box = results[0].boxes[0]
        predicted = results[0].names[int(box.cls)].strip().lower()
        confidence = float(box.conf)
        correct = predicted == expected[img_name]
        result = confidence if correct else "incorrect"
        display = f"{confidence:.0%}" if correct else "incorrect"
        print(f"{img_name:<40} {expected[img_name]:<12} {predicted:<12} {display:<12}")
        run_results[img_name] = result

# Add new column to tracker
run_label = datetime.now().strftime("%Y-%m-%d %H:%M")
next_col = ws.max_column + 1
ws.cell(row=1, column=next_col, value=run_label)

for row in range(2, ws.max_row + 1):
    img_name = ws.cell(row=row, column=1).value
    if img_name in run_results:
        cell = ws.cell(row=row, column=next_col, value=run_results[img_name])
        # Apply percent formatting only to numeric confidence values
        if isinstance(run_results[img_name], float):
            cell.number_format = '0%'

wb.save(tracker_path)
print(f"\nTracker updated: {tracker_path}")